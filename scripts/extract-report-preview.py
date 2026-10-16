import json
from pathlib import Path
import openpyxl

root = Path(__file__).resolve().parents[1]
sheet = openpyxl.load_workbook(root / 'public/templates/daily-report.xlsx').active
merged = { (m.min_row, m.min_col): m for m in sheet.merged_cells.ranges }
covered = { (r, c) for m in sheet.merged_cells.ranges for r in range(m.min_row, m.max_row + 1) for c in range(m.min_col, m.max_col + 1) if (r, c) != (m.min_row, m.min_col) }
def border(side):
    return ('2px' if side.style in ('medium', 'thick', 'double') else '1px') + ' solid #222' if side.style else 'none'
rows = []
for row in range(1, 43):
    cells = []
    for col in range(2, 33):
        if (row, col) in covered:
            continue
        cell = sheet.cell(row, col)
        merge = merged.get((row, col))
        right = sheet.cell(row, merge.max_col) if merge else cell
        bottom = sheet.cell(merge.max_row, col) if merge else cell
        style = {'fontSize': (cell.font.sz or 11) * 96 / 72, 'fontWeight': 'bold' if cell.font.b else 'normal', 'textAlign': cell.alignment.horizontal if cell.alignment.horizontal in ('left', 'center', 'right') else 'left', 'verticalAlign': {'center': 'middle'}.get(cell.alignment.vertical, cell.alignment.vertical or 'middle'), 'borderTop': border(cell.border.top), 'borderLeft': border(cell.border.left), 'borderRight': border(right.border.right), 'borderBottom': border(bottom.border.bottom)}
        cells.append({'address': cell.coordinate, 'value': cell.value or '', 'rowSpan': merge.max_row - row + 1 if merge else 1, 'colSpan': merge.max_col - col + 1 if merge else 1, 'style': style})
    rows.append({'height': (sheet.row_dimensions[row].height or sheet.sheet_format.defaultRowHeight or 15) * 96 / 72, 'cells': cells})
# The app's work rows use C:D for time and E:AE for subject plus surnames.
# Keep the original workbook unchanged; the Drive writer applies the same merges.
for number in range(14, 36):
    row = rows[number - 1]
    original = {cell['address']: cell for cell in row['cells']}
    time = dict(original['C' + str(number)])
    text = dict(original['D' + str(number)])
    time.update(address='C' + str(number), colSpan=2, value='')
    time['style'] = {**time['style'], 'borderRight': '1px solid #222', 'fontSize': 12, 'textAlign': 'center', 'verticalAlign': 'middle'}
    text.update(address='E' + str(number), colSpan=27, value='')
    text['style'] = {**text['style'], 'borderLeft': 'none', 'borderRight': original['AD' + str(number)]['style']['borderRight'], 'fontSize': 10 * 96 / 72, 'textAlign': 'left', 'verticalAlign': 'middle'}
    before = [cell for cell in row['cells'] if cell['address'].startswith('B')]
    after = [cell for cell in row['cells'] if cell['address'].startswith('AF')]
    row['cells'] = before + [time, text] + after
widths = []
for col in range(2, 33):
    dimension = next((d for d in sheet.column_dimensions.values() if d.min <= col <= d.max), None)
    widths.append(round((dimension.width if dimension else sheet.sheet_format.defaultColWidth or 8.43) * 7 + 5))
(root / 'src/dailyReportTemplate.json').write_text(json.dumps({'widths': widths, 'rows': rows}, ensure_ascii=False, indent=2), encoding='utf-8')
